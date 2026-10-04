import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
import docx
from pypdf import PdfReader


class Document:
    def __init__(self, content: str, metadata: Dict[str, Any]):
        self.content = content.strip()
        self.metadata = metadata

    def __repr__(self):
        return f"<Document source='{self.metadata.get('source_file')}' chars={len(self.content)}>"


class BankDocumentLoader:
    def __init__(self, faq_dir: str = "Amharic FAQ", text_dir: str = "amharic text"):
        self.faq_dir = Path(faq_dir)
        self.text_dir = Path(text_dir)

    def read_docx(self, file_path: Path) -> str:
        """Extract plain text from a .docx file."""
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            # Also extract tables if present
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        table_texts.append(row_text)
            return "\n".join(paragraphs + table_texts)
        except Exception as e:
            print(f"Warning: Failed to read docx {file_path}: {e}")
            return ""

    def read_pdf(self, file_path: Path) -> str:
        """Extract plain text from a .pdf file."""
        try:
            reader = PdfReader(file_path)
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(text.strip())
            return "\n".join(pages_text)
        except Exception as e:
            print(f"Warning: Failed to read pdf {file_path}: {e}")
            return ""

    def load_file(self, file_path: Path, category: str) -> Optional[Document]:
        """Load a single docx or pdf file."""
        ext = file_path.suffix.lower()
        content = ""
        if ext == ".docx":
            content = self.read_docx(file_path)
        elif ext == ".pdf":
            content = self.read_pdf(file_path)
        else:
            return None

        if not content or len(content) < 10:
            return None

        metadata = {
            "source_file": file_path.name,
            "category": category,
            "file_path": str(file_path),
            "file_type": ext[1:],
            "file_stem": file_path.stem
        }
        return Document(content=content, metadata=metadata)

    def load_all_documents(self) -> List[Document]:
        """Load all bank documents from FAQ and text directories."""
        documents = []

        # Load FAQ documents
        if self.faq_dir.exists():
            for f in sorted(self.faq_dir.iterdir()):
                if f.is_file() and f.suffix.lower() in [".docx", ".pdf"]:
                    doc = self.load_file(f, category="FAQ")
                    if doc:
                        documents.append(doc)

        # Load General Text documents
        if self.text_dir.exists():
            for f in sorted(self.text_dir.iterdir()):
                if f.is_file() and f.suffix.lower() in [".docx", ".pdf"]:
                    doc = self.load_file(f, category="General")
                    if doc:
                        documents.append(doc)

        print(f"Loaded total {len(documents)} raw documents ({sum(1 for d in documents if d.metadata['category']=='FAQ')} FAQ, {sum(1 for d in documents if d.metadata['category']=='General')} General).")
        return documents


if __name__ == "__main__":
    loader = BankDocumentLoader()
    docs = loader.load_all_documents()
    for d in docs[:5]:
        print(d, "Sample content snippet:", d.content[:100].replace("\n", " "))
